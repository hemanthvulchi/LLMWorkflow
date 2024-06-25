from pypdf import PdfReader
from pptx import Presentation
from docx import Document


def extract_text_from_pdf(file_path):
    text = ""
    reader = PdfReader(file_path)
    for page in reader.pages:
        text += page.extract_text()
    return text

def extract_text_from_pptx(file_path):
    text = ""
    prs = Presentation(file_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text

def extract_text_from_docx(file_path):
    text = ""
    doc = Document(file_path)
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text